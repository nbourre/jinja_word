import os
import sys

from docx import Document
import re
import csv


"""
Script permettant d'extraire les champs d'un Template Word respectant l'outil Jinja
"""

# Fichier source du template
# Changer pour le chemin du fichier à extraire les champs
fichier_source = r'D:\_data\projets\office\jinja_word\data\src_work.docx'

if not os.path.exists (fichier_source):
    print (f"Le fichier {fichier_source} est inexistant.")
    sys.exit()

# Permet d'extraire une liste de tous les champs du document
def extract_field_names(docx_path, placeholder_pattern=r"{{\s*([^|},]+)\s*}}"):
    doc = Document(docx_path)
    field_names = set()

    # Regular expression pattern to match placeholders
    pattern = re.compile(placeholder_pattern)

    for paragraph in doc.paragraphs:
        text = paragraph.text
        matches = pattern.findall(text)
        field_names.update(matches)

    return list(field_names)

# Permet de sauvegarder la liste dans un fichier de type CSV
# Ça s'ouvre dans Excel
def save_to_csv(template_path, field_names):
    csv_filename = template_path.replace('.docx', '_fields.csv')
    with open(csv_filename, 'w', newline='') as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(['Field Names'])
        writer.writerows([[field] for field in field_names])

field_names = extract_field_names(fichier_source)

print("Nom des champs extraits du template :")
for field_name in field_names:
    print(field_name)

# Sauvegarde du fichier avec les champs
save_to_csv(fichier_source, field_names)
